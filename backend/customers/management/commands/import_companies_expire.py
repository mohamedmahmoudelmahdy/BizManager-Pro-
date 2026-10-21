import datetime
from pathlib import Path

from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from customers.models import CompanyEmployee, Customer


def _clean_text(value):
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip().strip('"').strip()
    return str(value).strip()


def _as_date(value):
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    return None


def _as_int(value):
    try:
        return int(value)
    except Exception:
        return None


def _is_meaningful(value):
    if value in (None, "", ".", "-", "--"):
        return False
    return True


class Command(BaseCommand):
    help = "Import companies and their employees from Companies_expire.xlsx into the system."

    def add_arguments(self, parser):
        parser.add_argument(
            "--file",
            dest="file_path",
            default=None,
            help="Path to Companies_expire.xlsx. Defaults to project root /Companies_expire.xlsx",
        )
        parser.add_argument(
            "--limit",
            type=int,
            default=None,
            help="Optional limit of sheets to import (useful for testing).",
        )
        parser.add_argument(
            "--replace-employees",
            action="store_true",
            help="If set, existing employees for a company are deleted before importing.",
        )

    def handle(self, *args, **options):
        default_file = Path(settings.BASE_DIR).parent / "Companies_expire.xlsx"
        file_path = Path(options["file_path"] or default_file)
        if not file_path.exists():
            raise CommandError(f"File not found: {file_path}")

        workbook = load_workbook(file_path, data_only=True)
        sheet_names = [name for name in workbook.sheetnames if name.lower() not in {"compaines", "companies", "summary"}]
        if options.get("limit"):
            sheet_names = sheet_names[: options["limit"]]

        stats = {"companies_created": 0, "companies_updated": 0, "employees_created": 0, "employees_updated": 0, "skipped": 0}

        for sheet_name in sheet_names:
            sheet = workbook[sheet_name]
            company_name = _clean_text(sheet["A1"].value)
            as_of_date = _as_date(sheet["E1"].value)

            if not company_name:
                self.stdout.write(self.style.WARNING(f"[{sheet_name}] Skipped: missing company name"))
                stats["skipped"] += 1
                continue

            licenses, inferred_email, inferred_password = self._parse_licenses(sheet)
            employees = self._parse_employees(sheet)

            license_expiry = self._find_license_date(licenses, ["الرخصة", "license"])
            workers_expiry = self._find_license_date(licenses, ["بطاقة", "المنشاة", "المنشأة", "establishment"])
            health_insurance_expiry = self._find_license_date(licenses, ["الاشتراك", "annual", "subscription"])

            extra_data = {
                "source_sheet": sheet_name,
                "as_of_date": as_of_date.isoformat() if as_of_date else None,
                "licenses": [
                    {
                        "label": item["label"],
                        "expiry": item["expiry"],
                        "days_remaining": item["days_remaining"],
                        "extras": item["extras"],
                    }
                    for item in licenses
                ],
                "employees_raw": [
                    {
                        "name": emp["name"],
                        "residency_expiry": emp["residency_expiry"].isoformat() if emp["residency_expiry"] else None,
                        "residency_days": emp["residency_days"],
                        "guarantee_expiry": emp["guarantee_expiry"].isoformat() if emp["guarantee_expiry"] else None,
                        "guarantee_days": emp["guarantee_days"],
                    }
                    for emp in employees
                ],
            }
            credentials = {}
            if inferred_email:
                credentials["email"] = inferred_email
            if inferred_password:
                credentials["password"] = inferred_password
            if credentials:
                extra_data["credentials"] = credentials

            existing = Customer.objects.filter(name=company_name).first()
            defaults = {
                "type": "company",
                "name": company_name,
                "company_name": company_name,
                "phone": existing.phone if existing else "",
                "phone_secondary": existing.phone_secondary if existing else "",
                "email": inferred_email or (existing.email if existing else ""),
                "emirates_id": existing.emirates_id if existing else "",
                "trade_license": existing.trade_license if existing else "",
                "license_expiry": license_expiry or (existing.license_expiry if existing else None),
                "workers_expiry": workers_expiry or (existing.workers_expiry if existing else None),
                "health_insurance_expiry": health_insurance_expiry or (existing.health_insurance_expiry if existing else None),
                "visa_expiry": existing.visa_expiry if existing else None,
                "insurance_expiry": existing.insurance_expiry if existing else None,
                "balance": existing.balance if existing else 0,
                "debt_status": existing.debt_status if existing else "clear",
                "last_activity": as_of_date or (existing.last_activity if existing else timezone.localdate()),
                "extra_data": extra_data,
            }

            with transaction.atomic():
                customer, created = Customer.objects.update_or_create(
                    name=company_name,
                    defaults=defaults,
                )
                stats["companies_created" if created else "companies_updated"] += 1

                if options.get("replace_employees"):
                    customer.employees.all().delete()

                for emp in employees:
                    obj, emp_created = CompanyEmployee.objects.update_or_create(
                        company=customer,
                        name=emp["name"],
                        defaults={
                            "residency_expiry": emp["residency_expiry"],
                            "insurance_expiry": emp["guarantee_expiry"],
                        },
                    )
                    stats["employees_created" if emp_created else "employees_updated"] += 1

            self.stdout.write(
                self.style.SUCCESS(
                    f"[{sheet_name}] Imported company '{company_name}' "
                    f"with {len(employees)} employees; license expiry={license_expiry}, workers expiry={workers_expiry}"
                )
            )

        self.stdout.write(
            f"Done. Companies created: {stats['companies_created']}, updated: {stats['companies_updated']}, "
            f"employees created: {stats['employees_created']}, updated: {stats['employees_updated']}, skipped: {stats['skipped']}."
        )

    def _parse_licenses(self, sheet):
        licenses = []
        inferred_email = None
        inferred_password = None
        for row in sheet.iter_rows(min_row=3, values_only=True):
            label = _clean_text(row[0]) if row and len(row) > 0 else ""
            if not label:
                continue
            if label == "العمال":
                break
            expiry = _as_date(row[1]) if len(row) > 1 else None
            days_remaining = _as_int(row[2]) if len(row) > 2 else None
            extras = [_clean_text(value) for value in (row[3:] if len(row) > 3 else []) if _is_meaningful(value)]
            for val in extras:
                if not inferred_email and isinstance(val, str) and "@" in val:
                    inferred_email = val
                elif not inferred_password and isinstance(val, str) and any(ch.isdigit() for ch in val):
                    inferred_password = val
            licenses.append(
                {
                    "label": label,
                    "expiry": expiry.isoformat() if expiry else None,
                    "expiry_date": expiry,
                    "days_remaining": days_remaining,
                    "extras": extras,
                }
            )
        return licenses, inferred_email, inferred_password

    def _parse_employees(self, sheet):
        employees = []
        in_employees_section = False
        for row in sheet.iter_rows(values_only=True):
            name = _clean_text(row[0]) if row and len(row) > 0 else ""
            if name == "العمال":
                in_employees_section = True
                continue
            if not in_employees_section:
                continue
            if not name or name == ".":
                if all(value in (None, "", ".", "-", "--") for value in row[:5]):
                    break
                continue
            residency_expiry = _as_date(row[1]) if len(row) > 1 else None
            residency_days = _as_int(row[2]) if len(row) > 2 else None
            guarantee_expiry = _as_date(row[3]) if len(row) > 3 else None
            guarantee_days = _as_int(row[4]) if len(row) > 4 else None
            employees.append(
                {
                    "name": name,
                    "residency_expiry": residency_expiry,
                    "residency_days": residency_days,
                    "guarantee_expiry": guarantee_expiry,
                    "guarantee_days": guarantee_days,
                }
            )
        return employees

    @staticmethod
    def _find_license_date(licenses, keywords):
        for item in licenses:
            label_lower = item["label"].lower()
            if any(keyword in label_lower for keyword in keywords):
                return item["expiry_date"]
        return None
