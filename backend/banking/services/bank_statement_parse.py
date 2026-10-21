"""Parse uploaded bank statement files (CSV, XLSX, XLS) with header auto-detection."""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import BinaryIO

_COLUMN_KEYWORDS = {
    "date": ["date", "transaction date", "posting date", "value date", "تاريخ", "التاريخ"],
    "description": [
        "description", "desc", "details", "narrative", "memo", "particulars",
        "بيان", "تفاصيل", "الوصف", "البيان",
    ],
    "debit": ["debit", "debit amount", "withdrawal", "money out", "dr", "مدين", "سحب", "خصم"],
    "credit": ["credit", "credit amount", "deposit", "money in", "cr", "دائن", "إيداع", "ايداع"],
    "amount": ["amount", "value", "transaction amount", "مبلغ", "المبلغ", "قيمة"],
    "balance": ["balance", "running balance", "closing balance", "رصيد", "الرصيد"],
    "reference": [
        "reference", "ref", "bank reference", "customer reference", "transaction id",
        "مرجع", "رقم المرجع",
    ],
}

_REQUIRED_FIELDS = {"date", "description"}
_AMOUNT_FIELDS = {"debit", "credit", "amount"}


def _clean(value) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r\n", "\n").strip().strip('"').strip()
    return re.sub(r"\s+", " ", text)


def _norm_header(value) -> str:
    return _clean(value).casefold()


def _match_field(header: str) -> str | None:
    low = f" {_norm_header(header)} "
    for field_name, keywords in _COLUMN_KEYWORDS.items():
        for kw in keywords:
            if kw in low or _norm_header(header) == kw:
                return field_name
    return None


def _excel_serial_to_date(value: float) -> date | None:
    try:
        serial = float(value)
    except (TypeError, ValueError):
        return None
    if serial <= 0:
        return None
    base = datetime(1899, 12, 30)
    return (base + timedelta(days=serial)).date()


def parse_date_value(value) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    text = _clean(value)
    if not text or text in (".", "-", "--"):
        return None

    if re.fullmatch(r"\d+(\.\d+)?", text):
        parsed = _excel_serial_to_date(float(text))
        if parsed:
            return parsed

    for fmt in (
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%d-%b-%Y",
        "%d-%B-%Y",
        "%d %b %Y",
        "%d %B %Y",
        "%m/%d/%Y",
    ):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def parse_amount_value(value) -> Decimal | None:
    if value is None:
        return None
    if isinstance(value, (int, float, Decimal)):
        if value == 0:
            return None
        return Decimal(str(value)).quantize(Decimal("0.01"))

    text = _clean(value)
    if not text or text in (".", "-", "--"):
        return None
    text = text.replace(",", "").replace("AED", "").strip()
    if not text or text in (".", "-", "--"):
        return None
    try:
        amount = Decimal(text)
    except InvalidOperation:
        return None
    if amount == 0:
        return None
    return amount.quantize(Decimal("0.01"))


@dataclass
class ParsedBankRow:
    row_date: date | None
    description: str
    money_in: Decimal
    money_out: Decimal
    balance: Decimal | None
    reference: str
    raw: dict = field(default_factory=dict)


@dataclass
class ParseResult:
    rows: list[ParsedBankRow]
    detected_columns: dict[str, str]
    format_detected: str
    header_row_index: int


def _load_sheet_rows(file_obj: BinaryIO, filename: str) -> tuple[list[list[str]], str]:
    ext = Path(filename).suffix.lower()
    if ext == ".csv":
        raw = file_obj.read()
        if isinstance(raw, bytes):
            for encoding in ("utf-8-sig", "utf-8", "cp1256", "latin-1"):
                try:
                    text = raw.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue
            else:
                text = raw.decode("utf-8", errors="replace")
        else:
            text = raw
        reader = csv.reader(io.StringIO(text))
        return [list(row) for row in reader], "csv"

    if ext == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(file_obj, read_only=True, data_only=True)
        sheet = wb.active
        rows = []
        for row in sheet.iter_rows(values_only=True):
            rows.append([_clean(c) for c in row])
        wb.close()
        return rows, "xlsx"

    if ext == ".xls":
        import xlrd

        content = file_obj.read()
        wb = xlrd.open_workbook(file_contents=content)
        sheet = wb.sheet_by_index(0)
        rows = []
        for r in range(sheet.nrows):
            rows.append([_clean(sheet.cell_value(r, c)) for c in range(sheet.ncols)])
        return rows, "xls"

    raise ValueError("Unsupported file format. Upload CSV, XLSX, or XLS.")


def _detect_header_row(rows: list[list[str]]) -> tuple[int, dict[int, str]]:
    best_idx = -1
    best_map: dict[int, str] = {}
    best_score = 0

    for idx, row in enumerate(rows[:80]):
        col_map: dict[int, str] = {}
        for col_idx, cell in enumerate(row):
            field_name = _match_field(cell)
            if field_name and field_name not in col_map.values():
                col_map[col_idx] = field_name
        fields = set(col_map.values())
        has_amount = bool(fields & {"debit", "credit"})
        has_date = "date" in fields
        score = len(fields)
        if has_date:
            score += 2
        if has_amount:
            score += 2
        if score > best_score and has_date and has_amount:
            best_score = score
            best_idx = idx
            best_map = col_map

    if best_idx < 0:
        raise ValueError(
            "Could not detect bank statement headers. Expected columns for date, description, and debit/credit."
        )
    return best_idx, best_map


def _row_cells(row: list[str], col_map: dict[int, str]) -> dict[str, str]:
    out: dict[str, str] = {}
    for col_idx, field_name in col_map.items():
        if col_idx < len(row):
            out[field_name] = row[col_idx]
    return out


def _parse_data_row(cells: dict[str, str]) -> ParsedBankRow | None:
    row_date = parse_date_value(cells.get("date"))
    description = _clean(cells.get("description", ""))
    reference = _clean(cells.get("reference", ""))

    debit = parse_amount_value(cells.get("debit"))
    credit = parse_amount_value(cells.get("credit"))
    amount = parse_amount_value(cells.get("amount"))

    money_in = Decimal("0.00")
    money_out = Decimal("0.00")
    if credit:
        money_in = credit
    elif debit:
        money_out = debit
    elif amount:
        if amount < 0:
            money_out = abs(amount)
        else:
            money_in = amount

    balance = parse_amount_value(cells.get("balance"))

    if not row_date and not description and money_in == 0 and money_out == 0:
        return None
    if money_in == 0 and money_out == 0:
        return None

    return ParsedBankRow(
        row_date=row_date,
        description=description,
        money_in=money_in,
        money_out=money_out,
        balance=balance,
        reference=reference,
        raw=cells,
    )


def parse_bank_statement_file(file_obj: BinaryIO, filename: str) -> ParseResult:
    rows, fmt = _load_sheet_rows(file_obj, filename)
    if not rows:
        raise ValueError("The uploaded file is empty.")

    header_idx, col_map = _detect_header_row(rows)
    detected = {field: rows[header_idx][col] for col, field in col_map.items() if col < len(rows[header_idx])}

    parsed: list[ParsedBankRow] = []
    for row in rows[header_idx + 1 :]:
        if not any(_clean(c) for c in row):
            continue
        cells = _row_cells(row, col_map)
        item = _parse_data_row(cells)
        if item:
            parsed.append(item)

    if not parsed:
        raise ValueError("No transaction rows were found in the bank statement file.")

    return ParseResult(
        rows=parsed,
        detected_columns=detected,
        format_detected=fmt,
        header_row_index=header_idx,
    )
