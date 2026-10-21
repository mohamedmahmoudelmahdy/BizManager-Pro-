"""Workbook-specific sheet/column mapping for Excel import.

Extends the generic auto-detector (``import_excel``) with:
- per-sheet type + column mapping that the user can override,
- workbook profiles (generic / bizmanager_legacy),
- ledger-style sheets where the sheet/header is the account name and each row is
  a debit/credit ledger entry.

Nothing here posts to live balances — it only (re)generates preview rows.
"""
from __future__ import annotations

from pathlib import Path

from django.utils import timezone
from openpyxl import load_workbook

from financial_control.models import FinancialImportJob, FinancialImportRow
from financial_control.services.import_excel import _as_date, _as_decimal, _clean

# System fields a column can be mapped to.
SYSTEM_FIELDS = [
    "date",
    "description",
    "amount",
    "debit",
    "credit",
    "balance",
    "account_name",
    "payment_method",
    "check_number",
    "due_date",
    "payee",
    "notes",
]

# Sheet import types the user can pick.
SHEET_TYPES = [
    "skip",
    "ledger_account",
    "ledger_entries",
    "expense",
    "check",
    "debt",
    "transfer",
]

PROFILES = ["generic", "bizmanager_legacy", "custom"]

DT = FinancialImportRow.DetectedType
RS = FinancialImportRow.RowStatus

_TYPE_TO_DETECTED = {
    "ledger_account": DT.LEDGER_ACCOUNT,
    "ledger_entries": DT.LEDGER_ENTRY,
    "expense": DT.EXPENSE,
    "check": DT.CHECK,
    "debt": DT.DEBT,
    "transfer": DT.TRANSFER,
}

# Header keyword detection (English + Arabic).
_COLUMN_KEYWORDS = {
    "date": ["date", "تاريخ", "day", "yyyy"],
    "credit": ["credit", "دائن", "له", " cr", "cr "],
    "debit": ["debit", "مدين", "عليه", " dr", "dr "],
    "balance": ["balance", "رصيد", "الرصيد", "متبقي"],
    "amount": ["amount", "value", "مبلغ", "المبلغ", "قيمة", "total"],
    "description": [
        "description", "desc", "details", "narrative", "memo", "statement",
        "بيان", "تفاصيل", "ملاحظات", "وصف", "البيان",
    ],
    "check_number": ["check number", "check no", "cheque", "شيك", "رقم الشيك"],
    "due_date": ["due", "استحقاق", "تاريخ الاستحقاق"],
    "payee": ["payee", "vendor", "creditor", "party", "المستفيد", "المورد", "اسم", "العميل"],
    "payment_method": ["payment method", "paid by", "method", "طريقة الدفع", "الدفع", "bank"],
    "from_account": ["from", "from account", "source", "من"],
    "to_account": ["to", "to account", "destination", "إلى"],
}

# Arabic sheet-name terms → import type (payroll/employee deliberately mapped to skip).
_ARABIC_NAME_TERMS = [
    (("شيك", "شيكات", "check", "pdc"), "check"),
    (("دين", "ديون", "debt", "liabilit"), "debt"),
    (("تحويل", "تحويلات", "ايداع", "إيداع", "transfer", "deposit"), "transfer"),
    (("مصاريف", "مصروف", "expense"), "expense"),
    # payroll / employee — recognized but NOT imported yet
    (("راتب", "رواتب", "مرتب", "payroll", "salary", "سلفة", "مسحوبات"), "skip"),
]


def _match_field(text: str) -> str | None:
    low = f" {text.lower().strip()} "
    for field, kws in _COLUMN_KEYWORDS.items():
        for kw in kws:
            if kw in low or text.lower().strip() == kw.strip():
                return field
    return None


def _find_header(ws):
    """Return (header_row_index, [header cell values]) or (None, [])."""
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=15, values_only=True), start=1):
        cells = list(row)
        fields = set()
        for c in cells:
            txt = _clean(c)
            if not txt:
                continue
            f = _match_field(txt)
            if f:
                fields.add(f)
        # A usable header row has a date-ish column AND a money column.
        if ("date" in fields or "due_date" in fields) and (
            fields & {"credit", "debit", "amount", "balance"}
        ):
            return i, cells
    return None, []


def _map_columns(headers: list) -> dict:
    cols: dict[str, int] = {}
    for idx, cell in enumerate(headers):
        txt = _clean(cell)
        if not txt:
            continue
        field = _match_field(txt)
        if field and field not in cols:
            cols[field] = idx
    return cols


def _count_data_rows(ws, header_row: int) -> int:
    count = 0
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if any(_clean(c) for c in row):
            count += 1
    return count


def _suggest_type(sheet_name: str, cols: dict, profile: str) -> tuple[str, float]:
    name = (sheet_name or "").strip()
    low = name.lower()

    if profile == "bizmanager_legacy":
        if low in ("total", "summary", "اجمالي", "الاجمالي"):
            return "ledger_account", 0.7
        for terms, mapped in _ARABIC_NAME_TERMS:
            if any(term in low for term in terms):
                return mapped, 0.85
        if name.isdigit():
            if "date" in cols and (cols.get("credit") is not None or cols.get("debit") is not None):
                return "ledger_entries", 0.95
            return "ledger_entries", 0.7

    # Generic name-based hints.
    for terms, mapped in _ARABIC_NAME_TERMS:
        if any(term in low for term in terms) and mapped != "skip":
            return mapped, 0.6
    if any(k in low for k in ("ledger", "journal", "حساب")):
        return "ledger_entries", 0.6

    # Column-shape fallback: date + debit/credit looks like a ledger.
    if "date" in cols and (cols.get("credit") is not None or cols.get("debit") is not None):
        return "ledger_entries", 0.5
    if "date" in cols and "amount" in cols:
        return "expense", 0.35

    return "skip", 0.2


def analyze_sheet(ws, sheet_name: str, profile: str) -> dict:
    header_row, headers = _find_header(ws)
    if header_row is None:
        return {
            "sheet": sheet_name,
            "row_count": _count_data_rows(ws, 1),
            "header_row": 1,
            "columns": {},
            "header_labels": [],
            "account_name": sheet_name if sheet_name.strip().isdigit() else "",
            "detected_type": "skip",
            "confidence": 0.1,
        }

    cols = _map_columns(headers)
    account_name = ""
    # The "description" column in legacy ledger sheets carries the account name
    # as its header text (e.g. the vendor name). If description wasn't matched,
    # adopt the first unmapped non-empty header as description.
    used = set(cols.values())
    if "description" not in cols:
        for idx, h in enumerate(headers):
            txt = _clean(h)
            if txt and idx not in used and _match_field(txt) is None:
                cols["description"] = idx
                account_name = txt
                break
    if not account_name and "description" in cols:
        account_name = _clean(headers[cols["description"]])

    detected_type, confidence = _suggest_type(sheet_name, cols, profile)
    header_labels = [
        {"index": i, "label": _clean(h)} for i, h in enumerate(headers) if _clean(h)
    ]
    return {
        "sheet": sheet_name,
        "row_count": _count_data_rows(ws, header_row),
        "header_row": header_row,
        "columns": cols,
        "header_labels": header_labels,
        "account_name": account_name,
        "detected_type": detected_type,
        "confidence": confidence,
    }


def summarize_sheets(job: FinancialImportJob) -> dict:
    """Per-sheet analysis used by the mapping UI. Does not create rows."""
    path = Path(job.file.path)
    if not path.exists():
        return {"profile": job.profile or "generic", "sheets": [], "error": "File not found."}
    wb = load_workbook(path, read_only=True, data_only=True)
    profile = job.profile or "generic"
    saved = (job.mapping_json or {}).get("sheets", {})
    out = []
    for name in wb.sheetnames:
        analysis = analyze_sheet(wb[name], name, profile)
        user_cfg = saved.get(name) or {}
        analysis["user_type"] = user_cfg.get("type", "")
        analysis["user_columns"] = user_cfg.get("columns") or {}
        analysis["user_account_name"] = user_cfg.get("account_name", "")
        out.append(analysis)
    wb.close()
    return {"profile": profile, "fields": SYSTEM_FIELDS, "sheet_types": SHEET_TYPES, "sheets": out}


def _cell(cells, idx):
    if idx is None or idx < 0 or idx >= len(cells):
        return None
    return cells[idx]


def _ledger_row(cells, cols, account_name):
    """Build a normalized ledger-entry row and its validation result."""
    debit = _as_decimal(_cell(cells, cols.get("debit")))
    credit = _as_decimal(_cell(cells, cols.get("credit")))
    amount_col = _as_decimal(_cell(cells, cols.get("amount")))
    desc = _clean(_cell(cells, cols.get("description")))
    row_date = _as_date(_cell(cells, cols.get("date")))
    balance = _as_decimal(_cell(cells, cols.get("balance")))

    entry_type = "debit"
    amount = None
    warning = ""
    if debit and credit and debit > 0 and credit > 0:
        if debit != credit:
            warning = "Row has both debit and credit — review direction."
        amount, entry_type = debit, "debit"
    elif debit and debit > 0:
        amount, entry_type = debit, "debit"
    elif credit and credit > 0:
        amount, entry_type = credit, "credit"
    elif amount_col and amount_col > 0:
        amount, entry_type = amount_col, "debit"

    normalized = {
        "date": row_date.isoformat() if row_date else "",
        "due_date": "",
        "amount": float(amount) if amount else 0.0,
        "entry_type": entry_type,
        "description": desc,
        "payee": "",
        "account_name": account_name,
        "balance": float(balance) if balance is not None else None,
        "check_number": "",
        "type_hint": "",
        "from_account": "",
        "to_account": "",
        "payment_method": "",
    }
    if not amount or amount <= 0:
        return None, None, None  # empty / non-numeric row → skip entirely
    status = RS.WARNING if warning else RS.VALID
    return normalized, status, warning


def _generic_normalized(cells, cols):
    return {
        "date": (_as_date(_cell(cells, cols.get("date"))) or None),
        "due_date": (_as_date(_cell(cells, cols.get("due_date"))) or None),
        "amount": float(_as_decimal(_cell(cells, cols.get("amount"))) or 0),
        "description": _clean(_cell(cells, cols.get("description"))),
        "payee": _clean(_cell(cells, cols.get("payee"))),
        "check_number": _clean(_cell(cells, cols.get("check_number"))),
        "payment_method": _clean(_cell(cells, cols.get("payment_method"))),
        "from_account": _clean(_cell(cells, cols.get("from_account"))),
        "to_account": _clean(_cell(cells, cols.get("to_account"))),
        "type_hint": "",
    }


def validate_with_mapping(job: FinancialImportJob) -> FinancialImportJob:
    """Re-parse the workbook using the saved (or profile-suggested) mapping."""
    from financial_control.services.import_excel import _validate_row

    path = Path(job.file.path)
    if not path.exists():
        job.status = FinancialImportJob.Status.FAILED
        job.summary_json = {"error": "File not found on server."}
        job.save()
        return job

    wb = load_workbook(path, read_only=True, data_only=True)
    profile = job.profile or "generic"
    saved = (job.mapping_json or {}).get("sheets", {})

    job.rows.all().delete()
    total = valid = warning = error = 0
    by_type: dict = {}
    sheet_counts: dict = {}
    skipped_sheets = 0
    row_objects: list[FinancialImportRow] = []

    for name in wb.sheetnames:
        ws = wb[name]
        analysis = analyze_sheet(ws, name, profile)
        cfg = saved.get(name) or {}
        stype = cfg.get("type") or analysis["detected_type"]
        if stype == "skip":
            skipped_sheets += 1
            continue

        cols = {**analysis["columns"], **(cfg.get("columns") or {})}
        header_row = cfg.get("header_row") or analysis["header_row"]
        account_name = cfg.get("account_name") or analysis["account_name"] or name
        detected = _TYPE_TO_DETECTED.get(stype, DT.UNKNOWN)

        for row_idx, row in enumerate(
            ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1
        ):
            cells = list(row)
            if not any(_clean(c) for c in cells):
                continue

            if stype == "ledger_entries":
                normalized, status, warn = _ledger_row(cells, cols, account_name)
                if normalized is None:
                    continue
                err = ""
            elif stype == "ledger_account":
                acct = _clean(_cell(cells, cols.get("payee"))) or _clean(
                    _cell(cells, cols.get("description"))
                )
                if not acct:
                    continue
                normalized = {"account_name": acct, "amount": 0.0,
                              "balance": float(_as_decimal(_cell(cells, cols.get("balance"))) or 0)}
                status, warn, err = RS.VALID, "", ""
            else:
                normalized = _generic_normalized(cells, cols)
                normalized["date"] = (
                    normalized["date"].isoformat() if normalized["date"] else ""
                )
                normalized["due_date"] = (
                    normalized["due_date"].isoformat() if normalized["due_date"] else ""
                )
                status, warn, err = _validate_row(detected, normalized)

            total += 1
            if status == RS.VALID:
                valid += 1
            elif status == RS.WARNING:
                warning += 1
            else:
                error += 1

            bucket = by_type.setdefault(
                detected, {"count": 0, "valid": 0, "warning": 0, "error": 0, "total": 0.0}
            )
            bucket["count"] += 1
            bucket[status] = bucket.get(status, 0) + 1
            if status != RS.ERROR:
                bucket["total"] += float(normalized.get("amount") or 0)
            sheet_counts[name] = sheet_counts.get(name, 0) + 1

            row_objects.append(
                FinancialImportRow(
                    job=job,
                    sheet_name=name,
                    row_number=row_idx,
                    detected_type=detected,
                    raw_data_json={f"col_{i}": _clean(c) for i, c in enumerate(cells)},
                    normalized_data_json=normalized,
                    status=status,
                    mapping_status="mapped",
                    warning_message=warn,
                    error_message=err,
                )
            )

    wb.close()
    FinancialImportRow.objects.bulk_create(row_objects, batch_size=500)

    top_sheets = sorted(
        ({"name": k, "count": v} for k, v in sheet_counts.items()),
        key=lambda s: s["count"],
        reverse=True,
    )[:10]
    job.status = (
        FinancialImportJob.Status.VALIDATED if error == 0 else FinancialImportJob.Status.PREVIEW
    )
    job.total_rows = total
    job.valid_rows = valid
    job.warning_rows = warning
    job.error_rows = error
    job.completed_at = timezone.now()
    job.summary_json = {
        **(job.summary_json or {}),
        "confirm_disabled": True,
        "mapped": True,
        "skipped_sheets": skipped_sheets,
        "message": "Validated with mapping — confirm does not post until you review.",
        "preview_summary": {
            "by_type": by_type,
            "top_sheets": top_sheets,
            "unknown_rows": 0,
            "unknown_ratio": 0,
        },
    }
    job.save()
    return job
