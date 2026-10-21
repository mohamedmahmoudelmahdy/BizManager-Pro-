"""Excel import preview/validation — does not post to live balances."""
from __future__ import annotations

import hashlib
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.utils import timezone
from openpyxl import load_workbook

from financial_control.models import FinancialImportJob, FinancialImportRow


HEADER_ALIASES = {
    "date": {"date", "transaction date", "txn date", "تاريخ"},
    "amount": {"amount", "value", "debit", "credit", "total", "مبلغ", "المبلغ"},
    "description": {"description", "desc", "details", "narrative", "notes", "memo", "وصف"},
    "type": {"type", "category", "kind", "record type", "نوع"},
    "payee": {"payee", "vendor", "creditor", "name", "party", "employee", "المستفيد"},
    "check_number": {"check number", "check no", "cheque", "رقم الشيك"},
    "due_date": {"due date", "due", "تاريخ الاستحقاق"},
    "from_account": {"from", "from account", "source", "من"},
    "to_account": {"to", "to account", "destination", "إلى"},
    "payment_method": {
        "payment method",
        "paid by",
        "method",
        "account",
        "bank",
        "cash/bank",
        "bank/cash",
        "طريقة الدفع",
    },
}


def _clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def _as_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _clean(value)
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text[:10], fmt).date()
        except ValueError:
            continue
    return None


def _as_decimal(value) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    text = _clean(value).replace(",", "")
    text = re.sub(r"[^\d.\-]", "", text)
    if not text:
        return None
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def _map_headers(row_cells: list) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for idx, cell in enumerate(row_cells):
        key = _clean(cell).lower()
        if not key:
            continue
        for field, aliases in HEADER_ALIASES.items():
            if key in aliases or key == field:
                mapping[field] = idx
    return mapping


def _detect_type_from_sheet(sheet_name: str, row_type: str) -> str:
    name = sheet_name.lower()
    t = row_type.lower()
    combined = f"{name} {t}"
    if any(k in combined for k in ("check", "pdc", "cheque")):
        return FinancialImportRow.DetectedType.CHECK
    if any(k in combined for k in ("debt", "liability", "loan")):
        return FinancialImportRow.DetectedType.DEBT
    if any(k in combined for k in ("transfer", "deposit")):
        return FinancialImportRow.DetectedType.TRANSFER
    if any(k in combined for k in ("expense", "payment", "salary")):
        return FinancialImportRow.DetectedType.EXPENSE
    if any(k in combined for k in ("ledger", "journal")):
        return FinancialImportRow.DetectedType.LEDGER_ENTRY
    return FinancialImportRow.DetectedType.UNKNOWN


def _validate_row(detected: str, normalized: dict) -> tuple[str, str, str]:
    warning = ""
    error = ""
    amount = normalized.get("amount")

    if detected == FinancialImportRow.DetectedType.EXPENSE:
        if not normalized.get("date"):
            error = "Date is required for expenses."
        if amount is None or amount <= 0:
            error = error or "Amount must be greater than zero."
        if not normalized.get("payment_method"):
            error = error or "Payment method is required for expenses."
    elif detected == FinancialImportRow.DetectedType.CHECK:
        if not normalized.get("due_date"):
            error = "Due date is required for checks."
        if amount is None or amount <= 0:
            error = error or "Amount must be greater than zero."
        if not normalized.get("payee") and not normalized.get("check_number"):
            warning = "Payee or check number recommended."
    elif detected == FinancialImportRow.DetectedType.DEBT:
        if amount is None or amount <= 0:
            error = "Original amount is required for debts."
        if not normalized.get("payee") and not normalized.get("description"):
            warning = "Debt name/creditor recommended."
    elif detected == FinancialImportRow.DetectedType.TRANSFER:
        if amount is None or amount <= 0:
            error = "Transfer amount is required."
        if not normalized.get("from_account") or not normalized.get("to_account"):
            error = error or "From and to accounts are required."
    elif detected == FinancialImportRow.DetectedType.LEDGER_ENTRY:
        if amount is None or amount <= 0:
            error = "Ledger amount is required."
        if not normalized.get("description"):
            warning = "Description recommended."
    else:
        warning = "Could not detect row type — review manually."

    if error:
        status = FinancialImportRow.RowStatus.ERROR
    elif warning:
        status = FinancialImportRow.RowStatus.WARNING
    else:
        status = FinancialImportRow.RowStatus.VALID
    return status, warning, error


def _file_hash(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def run_import_preview(job: FinancialImportJob) -> FinancialImportJob:
    """Parse workbook, create preview rows, update job counters. No balance posting."""
    path = Path(job.file.path)
    if not path.exists():
        job.status = FinancialImportJob.Status.FAILED
        job.summary_json = {"error": "File not found on server."}
        job.save()
        return job

    job.rows.all().delete()
    digest = _file_hash(path)
    duplicate = (
        FinancialImportJob.objects.filter(duplicate_hash=digest)
        .exclude(pk=job.pk)
        .exists()
    )

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet_names = [
        n for n in workbook.sheetnames if n.lower() not in {"summary", "readme", "instructions"}
    ]

    total_rows = 0
    valid_rows = 0
    warning_rows = 0
    error_rows = 0
    unknown_rows = 0
    by_type: dict = {}
    sheet_counts: dict = {}
    row_objects: list[FinancialImportRow] = []

    for sheet_name in sheet_names:
        sheet = workbook[sheet_name]
        header_row = None
        header_map: dict[str, int] = {}
        for i, row in enumerate(sheet.iter_rows(min_row=1, max_row=15, values_only=True), start=1):
            cells = list(row)
            if not any(cells):
                continue
            trial = _map_headers(cells)
            if trial:
                header_row = i
                header_map = trial
                break
        if not header_map:
            continue

        for row_idx, row in enumerate(
            sheet.iter_rows(min_row=header_row + 1, values_only=True),
            start=header_row + 1,
        ):
            cells = list(row)
            if not any(_clean(c) for c in cells):
                continue
            raw = {f"col_{i}": _clean(c) for i, c in enumerate(cells)}
            normalized: dict = {}

            def cell(field: str):
                idx = header_map.get(field)
                if idx is None or idx >= len(cells):
                    return None
                return cells[idx]

            # Store dates as ISO strings so the JSONField (default encoder)
            # can serialize them. Raw date objects break the Postgres JSON write.
            row_date = _as_date(cell("date"))
            row_due = _as_date(cell("due_date")) or row_date
            normalized["date"] = row_date.isoformat() if row_date else ""
            normalized["due_date"] = row_due.isoformat() if row_due else ""
            normalized["amount"] = float(_as_decimal(cell("amount")) or 0)
            normalized["description"] = _clean(cell("description"))
            normalized["payee"] = _clean(cell("payee"))
            normalized["check_number"] = _clean(cell("check_number"))
            normalized["type_hint"] = _clean(cell("type"))
            normalized["from_account"] = _clean(cell("from_account"))
            normalized["to_account"] = _clean(cell("to_account"))
            normalized["payment_method"] = _clean(cell("payment_method"))

            detected = _detect_type_from_sheet(sheet_name, normalized.get("type_hint", ""))
            status, warning, error = _validate_row(detected, normalized)

            total_rows += 1
            if status == FinancialImportRow.RowStatus.VALID:
                valid_rows += 1
            elif status == FinancialImportRow.RowStatus.WARNING:
                warning_rows += 1
            else:
                error_rows += 1

            if detected == FinancialImportRow.DetectedType.UNKNOWN:
                unknown_rows += 1
            bucket = by_type.setdefault(
                detected, {"count": 0, "valid": 0, "warning": 0, "error": 0, "total": 0.0}
            )
            bucket["count"] += 1
            bucket[status] = bucket.get(status, 0) + 1
            if status != FinancialImportRow.RowStatus.ERROR:
                bucket["total"] += float(normalized.get("amount") or 0)
            sheet_counts[sheet_name] = sheet_counts.get(sheet_name, 0) + 1

            row_objects.append(
                FinancialImportRow(
                    job=job,
                    sheet_name=sheet_name,
                    row_number=row_idx,
                    detected_type=detected,
                    raw_data_json=raw,
                    normalized_data_json=normalized,
                    status=status,
                    warning_message=warning,
                    error_message=error,
                )
            )

    workbook.close()
    FinancialImportRow.objects.bulk_create(row_objects, batch_size=500)

    job.status = (
        FinancialImportJob.Status.VALIDATED
        if error_rows == 0
        else FinancialImportJob.Status.PREVIEW
    )
    job.total_sheets = len(sheet_names)
    job.total_rows = total_rows
    job.valid_rows = valid_rows
    job.warning_rows = warning_rows
    job.error_rows = error_rows
    job.duplicate_hash = digest
    job.completed_at = timezone.now()
    top_sheets = sorted(
        ({"name": k, "count": v} for k, v in sheet_counts.items()),
        key=lambda s: s["count"],
        reverse=True,
    )[:10]
    duplicate_job = (
        FinancialImportJob.objects.filter(duplicate_hash=digest)
        .exclude(pk=job.pk)
        .order_by("-created_at")
        .first()
    )
    job.summary_json = {
        "duplicate_file": duplicate,
        "duplicate_job_id": str(duplicate_job.id) if duplicate_job else None,
        "sheets": sheet_names,
        "confirm_disabled": True,
        "message": "Preview only — confirm does not post to live balances.",
        "preview_summary": {
            "by_type": by_type,
            "top_sheets": top_sheets,
            "unknown_rows": unknown_rows,
            "unknown_ratio": round(unknown_rows / total_rows, 3) if total_rows else 0,
        },
    }
    job.save()
    return job
